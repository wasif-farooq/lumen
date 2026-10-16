"""Pull plain text out of files so it can be full-text indexed (Strategy: one reader per type).

Readers yield the text piece by piece (a PDF page, a slide, a sheet row, a stretch of a text
file), so big files are never loaded whole; reading stops at limits.MAX_TEXT_CHARS."""

import codecs
import logging
from collections.abc import Callable, Iterable, Iterator
from pathlib import Path

from ..core import limits
from ..core.filetypes import TEXT_EXTS
from . import office

log = logging.getLogger(__name__)

Reader = Callable[[Path], Iterable[str]]


def extract_text(path: Path, ext: str, size: int) -> str:
    """Return the file's text, or "" when it isn't a readable document."""
    reader = extractor_for(ext)
    if reader is None or size == 0:
        return ""
    if ext in ZIPPED and size > limits.MAX_ZIPPED_BYTES:
        return ""  # unpacking these means reading the whole file
    try:
        return _collect(reader(path))
    except Exception as e:
        log.debug("no text from %s: %s", path, e)
        # Corrupt, encrypted or permission-denied files are indexed by name only
        return ""


def extractor_for(ext: str) -> Reader | None:
    """The reader for a file type, or None for types without text."""
    return EXTRACTORS.get(ext) or (_plain if ext in TEXT_EXTS else None)


def _collect(pieces: Iterable[str]) -> str:
    """Join the pieces, stopping once there is enough text (the reader is closed then)."""
    out, total = [], 0
    try:
        for piece in pieces:
            out.append(piece)
            total += len(piece) + 1
            if total >= limits.MAX_TEXT_CHARS:
                break
    except Exception as e:
        if not out:
            raise
        # A damaged end (or a part cut off by MAX_UNPACKED_BYTES): keep what was read
        log.debug("text cut short: %s", e)
    return "\n".join(out)[: limits.MAX_TEXT_CHARS]


def _plain(path: Path) -> Iterator[str]:
    decode = codecs.getincrementaldecoder("utf-8")(errors="ignore")
    with path.open("rb") as f:
        first = True
        while chunk := f.read(1 << 20):
            if first and b"\x00" in chunk[:4096]:
                return  # binary despite the extension
            first = False
            yield decode.decode(chunk)
        yield decode.decode(b"", final=True)


def _pdf(path: Path) -> Iterator[str]:
    import pymupdf

    with pymupdf.open(path) as doc:
        for page in doc:  # one page at a time: a long PDF is never loaded whole
            yield str(page.get_text())


def _docx(path: Path) -> Iterator[str]:
    import docx

    doc = docx.Document(str(path))
    for p in doc.paragraphs:
        yield p.text
    for table in doc.tables:
        for row in table.rows:
            yield " | ".join(cell.text for cell in row.cells)


# Readers for types that aren't plain text (TEXT_EXTS use _plain)
EXTRACTORS: dict[str, Reader] = {
    "pdf": _pdf,
    "docx": _docx,
    "xlsx": office.xlsx,
    "pptx": office.pptx,
    "odt": office.odf,
    "ods": office.odf,
    "odp": office.odf,
    "epub": office.epub,
}
ZIPPED = {"docx", "xlsx", "pptx", "odt", "ods", "odp", "epub"}  # zip files: size-capped
