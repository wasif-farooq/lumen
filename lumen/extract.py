"""Pull plain text out of files so it can be full-text indexed."""

from pathlib import Path
from .core import limits



TEXT_EXTS = {
    "txt", "md", "markdown", "rst", "csv", "tsv", "json", "jsonl", "yaml", "yml",
    "toml", "ini", "cfg", "conf", "log", "xml", "html", "htm", "tex", "srt",
    "py", "js", "ts", "tsx", "jsx", "mjs", "java", "kt", "c", "h", "cpp", "hpp",
    "cs", "go", "rs", "rb", "php", "swift", "dart", "sh", "bash", "sql", "css",
    "scss", "vue", "svelte", "ipynb", "gradle", "r", "lua",
}


def extract_text(path: Path, ext: str, size: int) -> str:
    """Return the file's text, or "" when it isn't a readable document."""
    if size == 0 or size > limits.MAX_EXTRACT_BYTES:
        return ""
    try:
        if ext == "pdf":
            return _pdf(path)
        if ext == "docx":
            return _docx(path)
        if ext in TEXT_EXTS:
            return _plain(path)
    except Exception:
        # Corrupt, encrypted or permission-denied files are indexed by name only
        return ""
    return ""


def _plain(path: Path) -> str:
    raw = path.read_bytes()[: limits.MAX_CONTENT_CHARS * 2]
    if b"\x00" in raw[:4096]:
        return ""  # binary despite the extension
    return raw.decode("utf-8", errors="ignore")[: limits.MAX_CONTENT_CHARS]


def _pdf(path: Path) -> str:
    import pymupdf

    parts, total = [], 0
    with pymupdf.open(path) as doc:
        for page in doc.pages(0, min(len(doc), limits.MAX_PDF_PAGES)):
            text = page.get_text()
            parts.append(text)
            total += len(text)
            if total >= limits.MAX_CONTENT_CHARS:
                break
    return "\n".join(parts)[: limits.MAX_CONTENT_CHARS]


def _docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)[: limits.MAX_CONTENT_CHARS]
