"""Documents read for search: office files and e-books, and big files read in pieces."""

import sqlite3

import pymupdf
import pytest
from conftest import index_file, index_now


def found(local, query: str) -> list[str]:
    return [h["name"] for h in local.get("/api/search", params={"q": query}).json()]


def text_of(local, path) -> str:
    return local.get("/api/file", params={"path": str(path)}).json()["excerpt"]


@pytest.fixture
def office_files(env):
    from docx import Document  # noqa: F401  (docx is covered elsewhere; these are the newer types)
    from ebooklib import epub
    from odf import opendocument, table, text
    from odf.text import H, P
    from openpyxl import Workbook
    from pptx import Presentation

    d = env.documents
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Travel to Lisbon", 1200])
    wb.create_sheet("Notes").append(["remember the parking receipts"])
    wb.save(d / "budget-2025.xlsx")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly roadmap"
    slide.placeholders[1].text = "Launch the harbour app in May"
    slide.notes_slide.notes_text_frame.text = "Mention the lighthouse partnership"
    prs.save(d / "roadmap.pptx")

    doc = opendocument.OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text="Garden plan"))
    doc.text.addElement(P(text="Plant tomatoes and basil along the south fence."))
    doc.save(str(d / "garden.odt"))

    sheet = opendocument.OpenDocumentSpreadsheet()
    t = table.Table(name="Books")
    for row in (["Title", "Author"], ["The Glass Orchard", "Mira Holt"]):
        tr = table.TableRow()
        for value in row:
            cell = table.TableCell()
            cell.addElement(text.P(text=value))
            tr.addElement(cell)
        t.addElement(tr)
    sheet.spreadsheet.addElement(t)
    sheet.save(str(d / "reading-list.ods"))

    book = epub.EpubBook()
    book.set_identifier("demo-1")
    book.set_title("The Quiet Harbour")
    book.set_language("en")
    chapters = [
        epub.EpubHtml(title="One", file_name="one.xhtml", content="<h1>One</h1><p>The ferry left at dawn.</p>"),
        epub.EpubHtml(title="Two", file_name="two.xhtml", content="<h1>Two</h1><p>A seagull stole the sandwich.</p>"),
    ]
    for c in chapters:
        book.add_item(c)
    book.toc = chapters
    book.spine = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(d / "quiet-harbour.epub"), book)
    return d


def test_office_documents_and_ebooks_are_searchable(office_files, local):
    index_now()
    assert found(local, "lisbon") == ["budget-2025.xlsx"]  # a cell
    assert found(local, "parking receipts") == ["budget-2025.xlsx"]  # on the second sheet
    assert found(local, "harbour app") == ["roadmap.pptx"]  # a slide
    assert found(local, "lighthouse partnership") == ["roadmap.pptx"]  # speaker notes
    assert found(local, "tomatoes basil") == ["garden.odt"]
    assert found(local, "glass orchard") == ["reading-list.ods"]
    assert found(local, "seagull sandwich") == ["quiet-harbour.epub"]  # the second chapter


def test_office_documents_keep_their_structure(office_files, local):
    index_now()
    sheet = text_of(local, office_files / "budget-2025.xlsx")
    assert "Sheet: Budget" in sheet and "Travel to Lisbon | 1200" in sheet and "Sheet: Notes" in sheet
    assert "The Glass Orchard | Mira Holt" in text_of(local, office_files / "reading-list.ods")
    book = text_of(local, office_files / "quiet-harbour.epub")
    assert book.index("ferry left at dawn") < book.index("seagull stole")  # reading order


def test_long_pdfs_are_read_to_the_end(env, local):
    pdf = pymupdf.open()
    for n in range(1, 151):
        pdf.new_page().insert_text((72, 72), f"Page {n} of the annual archive.")
    pdf[139].insert_text((72, 120), "The zeppelin hangar burned down in this year.")
    pdf.save(env.documents / "archive.pdf")
    index_now()
    assert found(local, "zeppelin hangar") == ["archive.pdf"]  # page 140: past the old 100-page limit


def test_text_past_the_old_limit_is_found_and_answers_questions(env, local, model):
    from conftest import tool_call

    notes = "\n".join(f"Day {i}: nothing unusual happened today." for i in range(12_000))  # ~480,000 characters
    (env.documents / "diary.txt").write_text(notes + "\nDay 12000: the octopus escaped from the aquarium.\n")
    index_now()
    assert found(local, "octopus aquarium") == ["diary.txt"]  # past the old 200,000-character limit

    model.script = [tool_call("find_in_files", question="when did the octopus escape?")]
    chat = local.post("/api/chats", json={}).json()["id"]
    events = local.post(f"/api/chats/{chat}/messages", json={"text": "when did the octopus escape?"}).text
    assert "diary.txt" in events
    tool_result = model.seen[-1][-1].content  # what the agent was given to answer from
    assert "Day 12000: the octopus escaped from the aquarium." in tool_result


def test_big_files_are_read_in_pieces_up_to_the_limit(env, local, text_limits):
    text_limits(MAX_TEXT_CHARS=100_000)
    big = env.documents / "server.log.txt"
    with big.open("w") as f:  # 30 MB: bigger than the old 25 MB cut-off, which skipped it whole
        f.write("startup: the walrus protocol is enabled\n")
        line = "request handled in 3 ms\n" * 1000
        for _ in range(1250):
            f.write(line)
        f.write("shutdown: the narwhal protocol is disabled\n")
    index_now()
    assert found(local, "walrus protocol") == ["server.log.txt"]  # the beginning is read
    assert found(local, "narwhal") == []  # the end is past the limit: not read
    assert len(text_of(local, big)) <= 100_000


def test_a_document_cut_short_keeps_what_was_read(env, local, text_limits):
    from openpyxl import Workbook

    wb = Workbook()
    for i in range(5000):
        wb.active.append([f"row {i}", "an ordinary expense"])
    wb.active.append(["the last row mentions a manatee"])
    wb.save(env.documents / "ledger.xlsx")
    text_limits(MAX_UNPACKED_BYTES=50_000)  # as a zip bomb guard would: the big sheet is cut off
    index_now()
    text = text_of(local, env.documents / "ledger.xlsx")
    assert "row 0 | an ordinary expense" in text  # the start is kept, not thrown away
    assert found(local, "manatee") == []


def test_an_index_from_0_2_reads_cut_short_and_new_document_types_again(office_files, local, text_limits):
    text_limits(MAX_TEXT_CHARS=200_000)  # as 0.2 read it
    (office_files / "long.txt").write_text("filler line\n" * 25_000 + "the platypus clause\n")
    index_now()
    assert found(local, "platypus clause") == []
    conn = sqlite3.connect(index_file())  # as 0.2 left it: no text for the new types, schema version 1
    conn.execute(
        "UPDATE files_fts SET content = '' WHERE rowid IN "
        "(SELECT id FROM files WHERE ext IN ('xlsx','pptx','odt','ods','odp','epub'))"
    )
    conn.execute("PRAGMA user_version = 1")
    conn.commit()
    conn.close()
    assert found(local, "lisbon") == []

    text_limits(MAX_TEXT_CHARS=10_000_000)
    index_now()  # opening the index upgrades it; the scan reads those files again
    assert found(local, "platypus clause") == ["long.txt"]
    assert found(local, "lisbon") == ["budget-2025.xlsx"]
